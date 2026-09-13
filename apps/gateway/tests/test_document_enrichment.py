from __future__ import annotations

import sqlite3
from contextlib import closing
from pathlib import Path


def _seed_document_database(path: Path, asset_root: Path) -> sqlite3.Connection:
    from wechat_ai.database import initialize_database

    asset_root.mkdir(parents=True, exist_ok=True)
    first = asset_root / "notes.txt"
    first.write_text("第一份会议记录", encoding="utf-8")
    connection = initialize_database(path)
    connection.execute(
        """
        INSERT INTO conversations(id, source_conversation_id, title, kind, message_count)
        VALUES ('conversation-1', 'conversation-1', 'Documents', 'direct', 2)
        """
    )
    for ordinal, file_name in enumerate(("notes.txt", "second.md")):
        cursor = connection.execute(
            """
            INSERT INTO messages(
                conversation_id, dedupe_key, sender_name, sent_at, type, content
            ) VALUES ('conversation-1', ?, 'Alice', ?, 'file', '')
            """,
            (f"document-{ordinal}", ordinal + 1),
        )
        part = connection.execute(
            """
            INSERT INTO message_parts(message_id, ordinal, kind, status)
            VALUES (?, 0, 'file', 'unresolved')
            """,
            (cursor.lastrowid,),
        )
        connection.execute(
            """
            INSERT INTO assets(
                part_id, ordinal, kind, status, file_name, mime_type,
                local_path, sha256
            ) VALUES (?, 0, 'file', 'resolved', ?, ?, ?, ?)
            """,
            (
                part.lastrowid,
                file_name,
                "text/plain" if file_name.endswith(".txt") else "text/markdown",
                file_name,
                "stale-hash",
            ),
        )
    connection.commit()
    return connection


def test_extract_document_text_reads_markdown_without_external_engine(
    tmp_path: Path,
) -> None:
    from wechat_ai.services.document_enrichment_service import extract_document_text

    document = tmp_path / "brief.md"
    document.write_text("# 项目简报\n\n结论：继续推进。", encoding="utf-8")

    assert extract_document_text(document) == "# 项目简报\n\n结论：继续推进。"


def test_extract_document_text_supports_docx_xlsx_and_pptx(
    tmp_path: Path,
) -> None:
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation

    from wechat_ai.services.document_enrichment_service import extract_document_text

    docx_path = tmp_path / "brief.docx"
    docx_document = Document()
    docx_document.add_paragraph("DOCX 会议结论")
    docx_document.save(docx_path)

    xlsx_path = tmp_path / "table.xlsx"
    workbook = Workbook()
    workbook.active.append(["项目", "状态"])
    workbook.active.append(["Alpha", "进行中"])
    workbook.save(xlsx_path)

    pptx_path = tmp_path / "slides.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "PPTX 关键结论"
    presentation.save(pptx_path)

    assert extract_document_text(docx_path) == "DOCX 会议结论"
    assert "Alpha\t进行中" in (extract_document_text(xlsx_path) or "")
    assert extract_document_text(pptx_path) == "PPTX 关键结论"


def test_extract_document_text_reads_searchable_pdf(tmp_path: Path) -> None:
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    from wechat_ai.services.document_enrichment_service import extract_document_text

    pdf_path = tmp_path / "report.pdf"
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {NameObject("/F1"): writer._add_object(font)}
            )
        }
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 72 720 Td (PDF text) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    with pdf_path.open("wb") as output:
        writer.write(output)

    assert extract_document_text(pdf_path) == "PDF text"


def test_document_enrichment_batches_and_persists_searchable_text(
    tmp_path: Path,
) -> None:
    from wechat_ai.services.document_enrichment_service import (
        DocumentEnrichmentService,
    )
    from wechat_ai.services.knowledge_projection_service import project_knowledge

    database_path = tmp_path / "library.sqlite"
    asset_root = tmp_path / "assets"
    second = asset_root / "second.md"
    asset_root.mkdir(parents=True, exist_ok=True)
    second.write_text("第二份会议记录", encoding="utf-8")

    with closing(_seed_document_database(database_path, asset_root)) as connection:
        service = DocumentEnrichmentService(asset_root=asset_root)
        first = service.enrich_pending(connection, limit=1)
        second_result = service.enrich_pending(connection, limit=1)
        project_knowledge(connection)
        indexed = connection.execute(
            "SELECT text_content FROM retrieval_chunks ORDER BY id"
        ).fetchall()

    assert first.ready == 1
    assert first.has_more is True
    assert second_result.ready == 1
    assert second_result.has_more is False
    assert [row[0] for row in indexed] == ["第一份会议记录", "第二份会议记录"]
