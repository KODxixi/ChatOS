from __future__ import annotations

import hashlib
import sqlite3
import time
from collections.abc import Callable
from dataclasses import dataclass
from pathlib import Path


DocumentTextExtractor = Callable[[Path, str | None, str | None], str | None]
_MAX_DOCUMENT_BYTES = 32 * 1024 * 1024
_MAX_DOCUMENT_CHARS = 200_000
_MAX_PDF_PAGES = 200
_MAX_SPREADSHEET_CELLS = 50_000


class DocumentEngineUnavailable(RuntimeError):
    pass


class UnsupportedDocumentType(ValueError):
    pass


@dataclass(frozen=True, slots=True)
class DocumentEnrichmentResult:
    ready: int = 0
    skipped: int = 0
    errors: int = 0
    attempted: int = 0
    no_text: int = 0
    engine_unavailable: int = 0
    has_more: bool = False


def extract_document_text(
    path: Path,
    *,
    file_name: str | None = None,
    mime_type: str | None = None,
    max_chars: int = _MAX_DOCUMENT_CHARS,
) -> str | None:
    """Extract bounded plain text from a supported local document."""

    if max_chars < 1:
        raise ValueError("max_chars must be positive")
    try:
        if path.stat().st_size > _MAX_DOCUMENT_BYTES:
            raise ValueError("document exceeds the size limit")
    except OSError as exc:
        raise ValueError("document is not readable") from exc

    document_format = _document_format(file_name or path.name, mime_type)
    if document_format in {"txt", "md", "csv"}:
        text = _read_text(path)
    elif document_format == "pdf":
        text = _read_pdf(path)
    elif document_format == "docx":
        text = _read_docx(path)
    elif document_format == "xlsx":
        text = _read_xlsx(path)
    elif document_format == "pptx":
        text = _read_pptx(path)
    else:
        raise UnsupportedDocumentType(document_format or "unknown")
    text = text.strip()
    return text[:max_chars] or None


class DocumentEnrichmentService:
    """Project local document bodies into the shared media derivative index."""

    def __init__(
        self,
        *,
        asset_root: Path,
        document_extractor: DocumentTextExtractor | None = None,
        clock: Callable[[], int] | None = None,
    ) -> None:
        self._asset_root = asset_root.resolve()
        self._document_extractor = document_extractor or (
            lambda path, file_name, mime_type: extract_document_text(
                path,
                file_name=file_name,
                mime_type=mime_type,
            )
        )
        self._clock = clock or (lambda: int(time.time() * 1000))

    def enrich_pending(
        self,
        connection: sqlite3.Connection,
        *,
        limit: int = 50,
    ) -> DocumentEnrichmentResult:
        if limit < 1:
            raise ValueError("limit must be positive")
        connection.row_factory = sqlite3.Row
        rows = connection.execute(
            """
            SELECT assets.id, assets.kind, assets.file_name, assets.mime_type,
                   assets.local_path, assets.sha256
            FROM assets
            WHERE assets.status = 'resolved'
              AND assets.kind IN ('file', 'attachment')
            ORDER BY assets.id ASC
            """
        ).fetchall()
        ready = skipped = errors = 0
        attempted = no_text = engine_unavailable = 0
        for row in rows:
            asset_id = int(row["id"])
            path = self._asset_path(row["local_path"])
            try:
                content_hash = _path_hash(path)
            except OSError:
                path = None
                content_hash = "missing"
            if path is not None and row["sha256"] != content_hash:
                connection.execute(
                    "UPDATE assets SET sha256 = ? WHERE id = ?",
                    (content_hash, asset_id),
                )
            existing = connection.execute(
                """
                SELECT status, error_code
                FROM media_derivatives
                WHERE asset_id = ? AND kind = 'document' AND content_hash = ?
                """,
                (asset_id, content_hash),
            ).fetchone()
            if existing is not None:
                if existing["status"] == "ready":
                    skipped += 1
                    continue
                if existing["status"] == "skipped" and existing["error_code"] in {
                    "NO_TEXT",
                    "UNSUPPORTED_FILE_TYPE",
                    "ASSET_MISSING",
                }:
                    skipped += 1
                    continue
            if attempted >= limit:
                break
            attempted += 1
            if path is None:
                self._store(
                    connection,
                    asset_id,
                    "skipped",
                    "",
                    "",
                    content_hash,
                    "ASSET_MISSING",
                )
                skipped += 1
                continue
            try:
                text = (
                    self._document_extractor(
                        path,
                        str(row["file_name"]) if row["file_name"] else None,
                        str(row["mime_type"]) if row["mime_type"] else None,
                    )
                    or ""
                ).strip()
                if not text:
                    self._store(
                        connection,
                        asset_id,
                        "skipped",
                        "",
                        "",
                        content_hash,
                        "NO_TEXT",
                    )
                    skipped += 1
                    no_text += 1
                    continue
                self._store(
                    connection,
                    asset_id,
                    "ready",
                    text,
                    "document",
                    content_hash,
                    None,
                )
                self._update_part_text(connection, asset_id, text)
                ready += 1
            except UnsupportedDocumentType:
                self._store(
                    connection,
                    asset_id,
                    "skipped",
                    "",
                    "",
                    content_hash,
                    "UNSUPPORTED_FILE_TYPE",
                )
                skipped += 1
            except DocumentEngineUnavailable:
                self._store(
                    connection,
                    asset_id,
                    "skipped",
                    "",
                    "",
                    content_hash,
                    "ENGINE_UNAVAILABLE",
                )
                skipped += 1
                engine_unavailable += 1
            except Exception:
                self._store(
                    connection,
                    asset_id,
                    "error",
                    "",
                    "",
                    content_hash,
                    "DOCUMENT_PARSE_FAILED",
                )
                errors += 1
        connection.commit()
        return DocumentEnrichmentResult(
            ready=ready,
            skipped=skipped,
            errors=errors,
            attempted=attempted,
            no_text=no_text,
            engine_unavailable=engine_unavailable,
            has_more=_has_pending_documents(connection),
        )

    def _asset_path(self, local_path: object) -> Path | None:
        if not isinstance(local_path, str) or not local_path:
            return None
        candidate = (self._asset_root / local_path).resolve()
        try:
            candidate.relative_to(self._asset_root)
        except ValueError:
            return None
        return candidate if candidate.is_file() else None

    def _store(
        self,
        connection: sqlite3.Connection,
        asset_id: int,
        status: str,
        text: str,
        engine: str,
        content_hash: str,
        error_code: str | None,
    ) -> None:
        now = self._clock()
        connection.execute(
            """
            INSERT INTO media_derivatives(
                asset_id, kind, status, text_content, engine,
                content_hash, error_code, created_at, updated_at
            ) VALUES (?, 'document', ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(asset_id, kind, content_hash) DO UPDATE SET
                status = excluded.status,
                text_content = excluded.text_content,
                engine = excluded.engine,
                error_code = excluded.error_code,
                updated_at = excluded.updated_at
            """,
            (
                asset_id,
                status,
                text,
                engine,
                content_hash,
                error_code,
                now,
                now,
            ),
        )

    @staticmethod
    def _update_part_text(
        connection: sqlite3.Connection,
        asset_id: int,
        text: str,
    ) -> None:
        connection.execute(
            """
            UPDATE message_parts
            SET text_content = ?, status = 'resolved'
            WHERE id = (SELECT part_id FROM assets WHERE id = ?)
            """,
            (text, asset_id),
        )


def _document_format(file_name: str, mime_type: str | None) -> str:
    mime = (mime_type or "").casefold()
    mime_formats = {
        "application/pdf": "pdf",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
        "application/vnd.openxmlformats-officedocument.presentationml.presentation": "pptx",
        "text/plain": "txt",
        "text/markdown": "md",
        "text/csv": "csv",
        "application/csv": "csv",
    }
    if mime in mime_formats:
        return mime_formats[mime]
    suffix = Path(file_name).suffix.casefold()
    if suffix in {".md", ".markdown"}:
        return "md"
    return suffix.removeprefix(".")


def _read_text(path: Path) -> str:
    payload = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "utf-16"):
        try:
            return payload.decode(encoding).replace("\r\n", "\n").replace("\r", "\n")
        except UnicodeDecodeError:
            continue
    return payload.decode("utf-8", errors="replace")


def _read_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise DocumentEngineUnavailable("pypdf is not installed") from exc
    reader = PdfReader(str(path))
    parts: list[str] = []
    for page in reader.pages[:_MAX_PDF_PAGES]:
        text = page.extract_text() or ""
        if text.strip():
            parts.append(text)
        if sum(len(part) for part in parts) >= _MAX_DOCUMENT_CHARS:
            break
    return "\n\n".join(parts)


def _read_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise DocumentEngineUnavailable("python-docx is not installed") from exc
    document = Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            if any(values):
                parts.append("\t".join(values))
    return "\n".join(parts)


def _read_xlsx(path: Path) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise DocumentEngineUnavailable("openpyxl is not installed") from exc
    workbook = load_workbook(str(path), read_only=True, data_only=True)
    parts: list[str] = []
    cells = 0
    try:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                values = ["" if value is None else str(value) for value in row]
                if any(value.strip() for value in values):
                    parts.append("\t".join(values).rstrip())
                cells += len(values)
                if cells >= _MAX_SPREADSHEET_CELLS:
                    return "\n".join(parts)
    finally:
        workbook.close()
    return "\n".join(parts)


def _read_pptx(path: Path) -> str:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise DocumentEngineUnavailable("python-pptx is not installed") from exc
    presentation = Presentation(str(path))
    parts: list[str] = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = shape.text.strip()
                if text:
                    parts.append(text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    values = [cell.text.strip() for cell in row.cells]
                    if any(values):
                        parts.append("\t".join(values))
    return "\n".join(parts)


def _has_pending_documents(connection: sqlite3.Connection) -> bool:
    row = connection.execute(
        """
        SELECT 1
        FROM assets
        WHERE status = 'resolved'
          AND kind IN ('file', 'attachment')
          AND NOT EXISTS (
                SELECT 1
                FROM media_derivatives
                WHERE media_derivatives.asset_id = assets.id
                  AND media_derivatives.kind = 'document'
                  AND media_derivatives.content_hash = COALESCE(assets.sha256, '')
                  AND (
                        media_derivatives.status = 'ready'
                        OR (
                            media_derivatives.status = 'skipped'
                            AND media_derivatives.error_code IN (
                                'NO_TEXT', 'UNSUPPORTED_FILE_TYPE', 'ASSET_MISSING'
                            )
                        )
                      )
            )
        LIMIT 1
        """
    ).fetchone()
    return row is not None


def _path_hash(path: Path | None) -> str:
    if path is None:
        return "missing"
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()
